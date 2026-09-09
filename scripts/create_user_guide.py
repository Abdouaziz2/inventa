from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Guide_utilisation_Inventa.docx"
ICON = ROOT / "public" / "inventa-icon.png"

NAVY = "0A1628"
BLUE = "1B3A6B"
GOLD = "C9972A"
LIGHT_GOLD = "F5D97A"
PALE_GOLD = "FFF8E7"
PALE_BLUE = "EEF3FA"
INK = "1C2430"
MUTED = "5F6B7A"
WHITE = "FFFFFF"
RED = "A61B1B"
GREEN = "217A4A"
LIGHT_GRAY = "F4F6F8"
LINE = "D9DEE5"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_font(run, size=11, color=INK, bold=False, italic=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_text(doc, text="", size=11, color=INK, bold=False, italic=False, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1 and getattr(doc, "_inventa_next_heading_page_break", False):
        p.paragraph_format.page_break_before = True
        doc._inventa_next_heading_page_break = False
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 11)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 5)
    run = p.add_run(text)
    set_font(
        run,
        size={1: 18, 2: 14, 3: 12}[level],
        color=NAVY if level == 1 else BLUE,
        bold=True,
    )
    return p


def add_step(doc, number, title, text):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.62)
    table.columns[1].width = Inches(5.88)
    prevent_row_split(table.rows[0])
    number_cell, text_cell = table.rows[0].cells
    number_cell.width = Inches(0.62)
    text_cell.width = Inches(5.88)
    set_cell_shading(number_cell, GOLD)
    set_cell_shading(text_cell, PALE_GOLD)
    for cell in (number_cell, text_cell):
        set_cell_margins(cell, top=100, bottom=100)
        set_cell_border(cell, color="E9D7A8")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = number_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(str(number)), size=15, color=WHITE, bold=True)
    p = text_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(title), size=11, color=NAVY, bold=True)
    p = text_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    set_font(p.add_run(text), size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_callout(doc, label, text, kind="info"):
    colors = {
        "info": (PALE_BLUE, BLUE),
        "warning": ("FFF1F0", RED),
        "success": ("EBF7F0", GREEN),
        "tip": (PALE_GOLD, "7A5A00"),
    }
    fill, accent = colors[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=180, end=180)
    set_cell_border(cell, color=accent, size="8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(label.upper()), size=9.5, color=accent, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text), size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        set_font(p.add_run(item), size=10.5, color=INK)


def add_quick_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, title, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_border(cell, color=NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(title), size=9.5, color=WHITE, bold=True)
    for row_values in rows:
        row = table.add_row()
        for index, (cell, value, width) in enumerate(zip(row.cells, row_values, widths)):
            cell.width = Inches(width)
            set_cell_shading(cell, WHITE if len(table.rows) % 2 else LIGHT_GRAY)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), size=9.5, color=INK, bold=index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_page_break(doc):
    doc._inventa_next_heading_page_break = True


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)

    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.7))
    table.autofit = False
    table.columns[0].width = Inches(3.35)
    table.columns[1].width = Inches(3.35)
    left, right = table.rows[0].cells
    left.width = right.width = Inches(3.35)
    for cell in (left, right):
        set_cell_margins(cell, top=0, bottom=40, start=0, end=0)
    p = left.paragraphs[0]
    set_font(p.add_run("INVENTA"), size=9, color=GOLD, bold=True)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("Guide d'utilisation"), size=9, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Inventa - Inventory Intelligence | Utilisation quotidienne"), size=8.5, color=MUTED)
    return doc


def build_document():
    doc = setup_document()

    # Cover
    add_text(doc, "GUIDE PRATIQUE", size=11, color=GOLD, bold=True, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    p.add_run().add_picture(str(ICON), width=Inches(1.55))
    add_text(doc, "Inventa", size=34, color=NAVY, bold=True, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(
        doc,
        "Guide d'utilisation pour les utilisateurs débutants",
        size=17,
        color=BLUE,
        bold=True,
        after=8,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(
        doc,
        "Gérer les clients, le stock, les dépôts, les réservations, les ventes et les reçus.",
        size=12,
        color=MUTED,
        after=22,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_callout(
        doc,
        "Objectif du guide",
        "Ce manuel explique les actions principales avec des mots simples. Suivez les étapes dans l'ordre et vérifiez toujours le récapitulatif avant de confirmer une opération.",
        "tip",
    )
    add_text(doc, "Version 1.0 - Juin 2026", size=9.5, color=MUTED, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_page_break(doc)

    # Overview
    add_heading(doc, "1. Découvrir Inventa", 1)
    add_text(
        doc,
        "Inventa est une application de gestion destinée aux boutiques. Elle permet de suivre les articles en stock, les clients et les opérations d'argent depuis un seul écran.",
    )
    add_heading(doc, "Les rubriques du menu", 2)
    add_quick_table(
        doc,
        ["Menu", "À quoi sert-il ?"],
        [
            ("Accueil", "Voir rapidement les ventes, dépôts, stocks et réservations."),
            ("Clients", "Ajouter un client et consulter son solde ou son historique."),
            ("Stock", "Ajouter, modifier et compter les bijoux ou articles."),
            ("Dépôt client", "Ajouter de l'argent sur le compte d'un client."),
            ("Réservations", "Bloquer un bijou pour un client avec un acompte."),
            ("Ventes", "Enregistrer une vente et encaisser le paiement."),
            ("Documents", "Retrouver, afficher et réimprimer les reçus et factures."),
            ("Paramètres", "Modifier le nom, l'adresse, les téléphones et le logo de la boutique."),
            ("Abonnement", "Consulter l'accès et les informations d'abonnement."),
        ],
        [1.35, 5.15],
    )
    add_callout(
        doc,
        "Bon à savoir",
        "Sur un téléphone ou un petit écran, appuyez sur le bouton Menu en haut à gauche pour afficher les rubriques.",
        "info",
    )

    add_heading(doc, "2. Installer et ouvrir l'application desktop", 1)
    add_step(doc, 1, "Ouvrir l'installateur", "Double-cliquez sur le fichier Inventa-Setup-1.0.0.exe.")
    add_step(doc, 2, "Autoriser l'installation", "Si Windows affiche une question de sécurité, vérifiez le nom Inventa puis continuez.")
    add_step(doc, 3, "Choisir le dossier", "Gardez le dossier proposé si vous ne savez pas lequel choisir.")
    add_step(doc, 4, "Terminer", "Laissez l'option de raccourci activée, puis cliquez sur Terminer.")
    add_step(doc, 5, "Lancer Inventa", "Double-cliquez sur l'icône Inventa présente sur le bureau ou dans le menu Démarrer.")
    add_callout(
        doc,
        "Attention",
        "N'installez l'application qu'à partir du fichier fourni par votre responsable. Ne partagez jamais votre mot de passe.",
        "warning",
    )

    add_page_break(doc)

    # Login and dashboard
    add_heading(doc, "3. Se connecter et se déconnecter", 1)
    add_heading(doc, "Se connecter", 2)
    add_step(doc, 1, "Saisir l'adresse email", "Cliquez dans la case Email et tapez l'adresse de votre compte.")
    add_step(doc, 2, "Saisir le mot de passe", "Cliquez dans la case Mot de passe. Les caractères restent cachés, c'est normal.")
    add_step(doc, 3, "Cliquer sur Se connecter", "Attendez l'ouverture du Tableau de bord.")
    add_heading(doc, "Se déconnecter", 2)
    add_text(
        doc,
        "Dans le bas du menu de gauche, cliquez sur la petite icône de sortie située à côté de votre nom.",
    )
    add_callout(
        doc,
        "En cas d'erreur",
        "Vérifiez l'adresse email, le mot de passe et la connexion Internet. Si l'accès reste impossible, contactez l'administrateur.",
        "warning",
    )

    add_heading(doc, "4. Comprendre le Tableau de bord", 1)
    add_bullets(
        doc,
        [
            "Total Ventes : montant total des ventes enregistrées.",
            "Total Dépôts : montant total ajouté sur les comptes clients.",
            "Valeur du Stock : valeur estimée des articles encore disponibles.",
            "Réservations : nombre de réservations enregistrées.",
            "Ventes par jour : graphique des ventes des derniers jours.",
            "Derniers dépôts : liste des dépôts les plus récents.",
        ],
    )
    add_callout(
        doc,
        "Conseil",
        "Commencez la journée par l'Accueil pour repérer un stock faible ou vérifier les dernières opérations.",
        "tip",
    )

    add_page_break(doc)

    # Clients
    add_heading(doc, "5. Gérer les clients", 1)
    add_heading(doc, "Ajouter un nouveau client", 2)
    add_step(doc, 1, "Ouvrir Clients", "Cliquez sur Clients dans le menu de gauche.")
    add_step(doc, 2, "Cliquer sur Nouveau Client", "Le formulaire apparaît dans une petite fenêtre.")
    add_step(doc, 3, "Remplir les informations", "Saisissez le nom complet et le numéro de téléphone.")
    add_step(doc, 4, "Cliquer sur Ajouter", "Un code client est créé automatiquement.")
    add_callout(
        doc,
        "Avant d'ajouter",
        "Utilisez la barre de recherche pour vérifier que le client n'existe pas déjà. Cela évite les doublons.",
        "warning",
    )
    add_heading(doc, "Rechercher et consulter un client", 2)
    add_step(doc, 1, "Taper un mot", "Recherchez avec le nom, le code client ou le téléphone.")
    add_step(doc, 2, "Ouvrir la fiche", "Cliquez sur Détails ou directement sur le client depuis un téléphone.")
    add_step(
        doc,
        3,
        "Lire la fiche",
        "Vous pouvez voir le solde disponible, l'historique des dépôts, les achats et les mouvements du portefeuille.",
    )
    add_callout(
        doc,
        "Comprendre le solde",
        "Le solde est l'argent disponible sur le compte du client. Il peut être utilisé pendant une vente lorsque l'option Utiliser le solde est cochée.",
        "info",
    )

    # Stock
    add_heading(doc, "6. Gérer le stock", 1)
    add_heading(doc, "Ajouter un bijou ou un article", 2)
    add_step(doc, 1, "Ouvrir Stock", "Cliquez sur Stock, puis sur + Ajouter un bijou.")
    add_step(doc, 2, "Indiquer le nom", "Exemple : Bague solitaire, Bracelet homme ou Collier argent.")
    add_step(doc, 3, "Choisir la matière et la catégorie", "Sélectionnez les valeurs dans les listes.")
    add_step(doc, 4, "Saisir le stock et le poids", "Le stock est la quantité disponible. Le poids est saisi en grammes.")
    add_step(
        doc,
        5,
        "Saisir les prix",
        "Indiquez le prix d'achat et le prix unitaire par gramme. Inventa calcule le prix total de vente et la marge estimée.",
    )
    add_step(doc, 6, "Ajouter une photo", "Cette étape est facultative. Cliquez sur Importer une image.")
    add_step(doc, 7, "Enregistrer", "Relisez les informations puis cliquez sur Enregistrer.")
    add_callout(
        doc,
        "Vérification importante",
        "Contrôlez le poids et le prix par gramme avant d'enregistrer. Ces deux valeurs déterminent le prix total de vente.",
        "warning",
    )
    add_heading(doc, "Modifier une fiche ou ajuster le stock", 2)
    add_step(doc, 1, "Trouver l'article", "Utilisez la recherche, le filtre de statut ou le tri.")
    add_step(doc, 2, "Ouvrir les actions", "Cliquez sur les trois points à droite de l'article.")
    add_step(doc, 3, "Choisir l'action", "Sélectionnez Modifier la fiche, Entrée de stock ou Sortie de stock.")
    add_step(doc, 4, "Indiquer un motif", "Pour une entrée ou une sortie, écrivez la raison du mouvement.")
    add_callout(
        doc,
        "Exemples de motif",
        "Réception fournisseur, correction d'inventaire, article cassé, perte ou retour en stock.",
        "tip",
    )

    add_page_break(doc)

    # Deposits
    add_heading(doc, "7. Enregistrer un dépôt client", 1)
    add_text(
        doc,
        "Un dépôt ajoute de l'argent au compte d'un client. Le client pourra utiliser ce solde lors d'une vente.",
    )
    add_step(doc, 1, "Ouvrir Dépôt client", "Cliquez sur Dépôt client dans le menu.")
    add_step(doc, 2, "Rechercher le client", "Tapez son nom, son code ou son téléphone, puis cliquez sur le bon résultat.")
    add_step(doc, 3, "Saisir le montant", "Tapez le montant ou utilisez un bouton de montant rapide.")
    add_step(doc, 4, "Ajouter une note", "Ouvrez Plus d'options si vous souhaitez préciser la raison du dépôt.")
    add_step(
        doc,
        5,
        "Vérifier le récapitulatif",
        "Contrôlez le solde actuel, le dépôt ajouté et le nouveau solde.",
    )
    add_step(doc, 6, "Confirmer", "Cliquez sur Confirmer le dépôt.")
    add_step(doc, 7, "Imprimer le reçu", "À l'ouverture du reçu, cliquez sur Imprimer ou Ouvrir version PDF.")
    add_callout(
        doc,
        "Montant important",
        "Pour un dépôt élevé, Inventa demande une deuxième confirmation. Relisez attentivement le montant avant de continuer.",
        "warning",
    )

    add_page_break(doc)

    # Reservations
    add_heading(doc, "8. Réserver un bijou", 1)
    add_step(doc, 1, "Ouvrir Réservations", "Cliquez sur Réservations dans le menu.")
    add_step(doc, 2, "Choisir le client", "Sélectionnez le client dans la première liste.")
    add_step(doc, 3, "Choisir le bijou", "Seuls les bijoux disponibles et en stock sont proposés.")
    add_step(doc, 4, "Saisir l'acompte", "Le montant doit être supérieur à zéro et ne doit pas dépasser le prix du bijou.")
    add_step(doc, 5, "Choisir la date limite", "Indiquez la date avant laquelle la réservation doit être traitée.")
    add_step(
        doc,
        6,
        "Vérifier le reste à payer",
        "Inventa calcule automatiquement le prix total, l'acompte et le reste à payer.",
    )
    add_step(doc, 7, "Confirmer", "Cliquez sur Confirmer la Réservation puis imprimez le bon si nécessaire.")
    add_heading(doc, "Annuler une réservation", 2)
    add_text(
        doc,
        "Dans le suivi des réservations, cliquez sur Annuler pour une réservation active. Confirmez l'annulation. Une unité est remise automatiquement dans le stock.",
    )
    add_callout(
        doc,
        "Attention",
        "Vérifiez le numéro du document et le nom du client avant d'annuler une réservation.",
        "warning",
    )

    add_page_break(doc)

    # Sales
    add_heading(doc, "9. Enregistrer une vente", 1)
    add_heading(doc, "Préparer la vente", 2)
    add_step(doc, 1, "Ouvrir Ventes", "Cliquez sur Ventes dans le menu.")
    add_step(
        doc,
        2,
        "Choisir le client",
        "Sélectionnez un client existant. Vous pouvez aussi choisir Nouveau client et saisir son nom et son téléphone.",
    )
    add_step(
        doc,
        3,
        "Décider pour le solde",
        "Pour un client existant, laissez Utiliser le solde coché si son argent disponible doit payer une partie de la vente.",
    )
    add_step(doc, 4, "Ajouter les articles", "Choisissez un bijou, indiquez la quantité puis cliquez sur Ajouter.")
    add_step(doc, 5, "Contrôler le panier", "Vérifiez chaque article, la quantité et le total de la facture.")
    add_heading(doc, "Encaisser", 2)
    add_step(doc, 6, "Saisir le montant remis", "Tapez l'argent donné par le client.")
    add_step(doc, 7, "Choisir le mode de paiement", "Exemples : Espèces, Mobile Money, Carte ou Virement bancaire.")
    add_step(
        doc,
        8,
        "Lire le résumé",
        "Vérifiez le total, le solde utilisé, le montant payé, le reste à payer et la monnaie à rendre.",
    )
    add_step(doc, 9, "Confirmer la vente", "Cliquez sur Encaisser ou sur Enregistrer avec un reste à payer.")
    add_step(doc, 10, "Remettre la facture", "Imprimez la facture ou ouvrez sa version PDF.")
    add_callout(
        doc,
        "Vente à crédit",
        "Si un reste à payer existe, cochez l'autorisation de vente à crédit. Inventa demandera une confirmation supplémentaire.",
        "warning",
    )
    add_callout(
        doc,
        "Monnaie à rendre",
        "Quand le client remet plus que le montant demandé, la monnaie à rendre apparaît dans le résumé.",
        "info",
    )

    add_page_break(doc)

    # Documents
    add_heading(doc, "10. Retrouver et réimprimer un reçu", 1)
    add_step(doc, 1, "Ouvrir Documents", "Cliquez sur Documents dans le menu.")
    add_step(
        doc,
        2,
        "Repérer le document",
        "La liste affiche le type, le numéro, le client, le paiement, la date et le montant.",
    )
    add_step(doc, 3, "Ouvrir le reçu", "Cliquez sur la ligne du document ou sur Voir.")
    add_step(doc, 4, "Contrôler les informations", "Vérifiez le client, les articles, les montants et la date.")
    add_step(doc, 5, "Réimprimer", "Cliquez sur Imprimer. Vous pouvez aussi choisir Ouvrir version PDF.")
    add_callout(
        doc,
        "Conseil d'impression",
        "Dans les options d'impression du navigateur, désactivez En-têtes et pieds de page pour obtenir une facture propre.",
        "tip",
    )
    add_heading(doc, "Les trois types de documents", 2)
    add_quick_table(
        doc,
        ["Type", "Utilisation"],
        [
            ("Reçu de dépôt", "Confirme l'argent ajouté sur le compte du client."),
            ("Bon de réservation", "Confirme le bijou réservé, l'acompte et le reste à payer."),
            ("Facture de vente", "Confirme les articles vendus et les informations de paiement."),
        ],
        [1.75, 4.75],
    )

    add_page_break(doc)

    # Settings and good practices
    add_heading(doc, "11. Modifier les informations de la boutique", 1)
    add_step(doc, 1, "Ouvrir Paramètres", "Cliquez sur Paramètres dans le menu.")
    add_step(doc, 2, "Changer le logo", "Cliquez sur Changer le logo puis choisissez une image PNG, JPG, JPEG ou WEBP.")
    add_step(
        doc,
        3,
        "Remplir le profil",
        "Vérifiez le nom complet, le nom de la boutique, les téléphones et l'adresse.",
    )
    add_step(doc, 4, "Enregistrer", "Cliquez sur Enregistrer.")
    add_callout(
        doc,
        "Pourquoi c'est important",
        "Le nom, le logo, l'adresse et les téléphones de la boutique apparaissent sur les factures et les reçus.",
        "info",
    )

    add_heading(doc, "12. Bonnes pratiques quotidiennes", 1)
    add_bullets(
        doc,
        [
            "Vérifiez le nom du client avant chaque dépôt, réservation ou vente.",
            "Relisez toujours les montants avant de confirmer.",
            "Utilisez la recherche avant de créer un nouveau client.",
            "Indiquez un motif clair pour chaque entrée ou sortie de stock.",
            "Ne communiquez jamais votre mot de passe.",
            "Déconnectez-vous lorsque vous quittez le poste.",
            "Conservez les factures importantes en PDF ou sur papier selon les règles de la boutique.",
            "Signalez immédiatement toute opération incorrecte à l'administrateur.",
        ],
    )

    add_page_break(doc)

    # Troubleshooting and glossary
    add_heading(doc, "13. Aide en cas de problème", 1)
    add_quick_table(
        doc,
        ["Problème", "Solution simple"],
        [
            ("Connexion impossible", "Vérifiez Internet, l'email et le mot de passe. Contactez l'administrateur si nécessaire."),
            ("Bouton désactivé", "Un champ obligatoire manque ou contient une valeur incorrecte."),
            ("Client introuvable", "Essayez le nom, le code ou le téléphone. Vérifiez aussi l'orthographe."),
            ("Bijou absent de la vente", "Vérifiez qu'il est disponible et que son stock est supérieur à zéro."),
            ("Impression bloquée", "Autorisez les fenêtres contextuelles, puis réessayez depuis Documents."),
            ("Montant incorrect", "N'enregistrez pas l'opération. Revenez au champ et corrigez la valeur."),
            ("Application figée", "Fermez puis rouvrez Inventa. Ne confirmez pas deux fois la même opération."),
        ],
        [2.05, 4.45],
    )
    add_heading(doc, "Petit lexique", 2)
    add_quick_table(
        doc,
        ["Mot", "Signification"],
        [
            ("Solde client", "Argent disponible sur le compte du client."),
            ("Dépôt", "Argent ajouté au solde du client."),
            ("Acompte", "Première somme versée pour une réservation."),
            ("Reste à payer", "Montant que le client doit encore régler."),
            ("Stock", "Quantité disponible d'un article."),
            ("Entrée de stock", "Quantité ajoutée au stock."),
            ("Sortie de stock", "Quantité retirée du stock."),
            ("FCFA", "Franc CFA, monnaie utilisée pour les montants."),
        ],
        [1.65, 4.85],
    )
    add_callout(
        doc,
        "Règle principale",
        "En cas de doute, ne confirmez pas l'opération. Demandez une vérification à votre responsable.",
        "success",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
